"""Word 文档生成的公共封装：统一字体、标题层级、表格样式。

中文字体必须同时设置 ascii 和 eastAsia 两个属性，否则 Word 里会显示成方框或宋体，
这是 python-docx 最常见的坑。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

CN_FONT = "微软雅黑"
EN_FONT = "Segoe UI"
FOOTER_TEXT = "轻岚优选 · 知识库素材（合成数据，仅用于内部检索验证）"


def _cn(run, size: float = 10.5, bold: bool = False) -> None:
    run.font.name = EN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), CN_FONT)


def new_doc() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = EN_FONT
    style.font.size = Pt(10.5)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CN_FONT)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        footer = sec.footer.paragraphs[0]
        footer.text = FOOTER_TEXT
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            _cn(run, 8)
    return doc


def doc_title(doc: Document, text: str, sub: str = "") -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _cn(p.add_run(text), 20, bold=True)
    if sub:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _cn(p2.add_run(sub), 11)


def heading(doc: Document, text: str, level: int = 1) -> None:
    sizes = {1: 13.5, 2: 12, 3: 11}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    _cn(p.add_run(text), sizes.get(level, 11), bold=True)


def para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    _cn(p.add_run(text))


def numbered(doc: Document, items: list[str]) -> None:
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(2)
        _cn(p.add_run(f"{i}. {it}"))


def bullets(doc: Document, items: list[str]) -> None:
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(2)
        _cn(p.add_run(f"· {it}"))


def kv_table(doc: Document, pairs: list[tuple[str, str]]) -> None:
    """两列信息表，用于展示商品参数、订单信息这类键值对。"""
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in pairs:
        cells = t.add_row().cells
        cells[0].text = ""
        cells[1].text = ""
        _cn(cells[0].paragraphs[0].add_run(str(k)), 10, bold=True)
        _cn(cells[1].paragraphs[0].add_run(str(v)), 10)
    doc.add_paragraph()


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        _cn(cell.paragraphs[0].add_run(h), 10, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            _cn(cells[i].paragraphs[0].add_run(str(v)), 10)
    doc.add_paragraph()


def save(doc: Document, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
