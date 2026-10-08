"""把 500 条 FAQ 输出成规范 Excel（含常见问法变体，可直接入库）。

    python gen/faq_excel.py

三个 sheet：FAQ 清单（500 条）、分类统计、入库说明。
"""

from __future__ import annotations

import csv
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

DATA = HERE.parent / "data"
OUT = HERE.parent / "output" / "04-FAQ清单"

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)


def style_header(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> int:
    rows = list(csv.DictReader((DATA / "qa_seed.csv").open(encoding="utf-8-sig")))
    wb = Workbook()

    ws = wb.active
    ws.title = "FAQ清单"
    ws.append(["QA编号", "分类", "标准问题", "常见问法（可直接作为检索入口）",
               "标准答案", "适用品类", "关联政策码", "状态"])
    for r in rows:
        ws.append([r["qa_id"], r["分类"], r["问题"], r["常见问法"],
                   r["标准答案"], r["适用品类"], r["关联政策码"], "待入库"])
    style_header(ws, [10, 8, 30, 60, 70, 12, 26, 8])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    ws2 = wb.create_sheet("分类统计")
    ws2.append(["分类", "条数", "占比"])
    dist = Counter(r["分类"] for r in rows)
    for k, v in dist.most_common():
        ws2.append([k, v, f"{v / len(rows):.1%}"])
    ws2.append(["合计", len(rows), "100.0%"])
    style_header(ws2, [12, 10, 10])

    ws3 = wb.create_sheet("入库说明")
    notes = [
        ["项目", "说明"],
        ["用途", "FAQ 清单，补全常见问法后直接入 Dify 知识库"],
        ["建议切块", "一条 QA 为一个块，不要跨条合并，避免答案串味"],
        ["检索入口", "「常见问法」列已包含 5 种口语表达，入库时作为同一条 QA 的多种问法"],
        ["分批建议", "先入 50 条验证召回，再全量入库"],
        ["数据性质", "合成数据，仅用于内部检索与引用链路验证"],
        ["维护节奏", "每周根据人工工单的未解决原因补充新条目"],
    ]
    for row in notes:
        ws3.append(row)
    style_header(ws3, [14, 80])

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "FAQ清单_500条.xlsx"
    wb.save(path)
    print(f"完成：{path}")
    print(f"  FAQ {len(rows)} 条，分布 {dict(dist)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
